"""Read historical source workbooks; prepare private review drafts, never modify sources."""
import hashlib
import json
import sys
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'backups/spending-history-review-20261019'
labels = ['ค่ายาและเวชภัณฑ์','ค่าวัสดุ','ค่าตอบแทนทางการแพทย์','ค่าบริการทางการแพทย์','ค่าครุภัณฑ์ ที่ดินและสิ่งก่อสร้าง','ค่าใช้สอย','ค่าสาธารณูปโภค','ค่าจ้างลูกจ้างชั่วคราว','ค่าตอบแทนการปฏิบัติงานนอกเวลาราชการ','ค่าใช้จ่ายในการเดินทางไปราชการ','ค่าใช้จ่ายอื่นที่เกี่ยวข้องกับการสาธารณสุข']
assessment = json.loads((OUT / 'candidates.json').read_text(encoding='utf8'))
num = lambda v: round(float(v), 2) if isinstance(v,(int,float)) and not isinstance(v,bool) else None
decimal_number = lambda v: None if v is None else float(v)
records = []
for r in assessment['candidate_records']:
    if r['kind'] != 'actual':
        continue
    src = assessment['sources'][r['source_id']]
    assert hashlib.sha256(Path(src['path']).read_bytes()).hexdigest() == src['sha256'], 'Source changed; reassess first'
    groups = []
    for c in r['categories']:
        details = [{'id':'r'+d['cell'][1:], 'label':str(d['label'])[:500], 'amount':decimal_number(d['value']), 'commitment':None, 'planned':None}
                   for d in c['details'] if d['label'] is not None]
        groups.append({'id':c['id'], 'label':labels[int(c['id'])-1], 'amount':decimal_number(c['candidate_amount']), 'commitment':None, 'planned':None, 'note':'', 'details':details})
    plans = [x for x in assessment['candidate_records'] if x['kind']=='plan' and x['fiscal_year']==r['fiscal_year'] and x['facility_code']==r['facility_code']]
    if plans:
        original = next(x for x in plans if x['source_id']==2)
        for g, p in zip(groups,original['categories']):
            g['planned'] = decimal_number(p['reported_amount'])
    issues = ['ยังไม่ยืนยันวันเริ่มและวันสิ้นสุดของยอดจ่ายจริง ไม่ถือเป็นยอดปิดปีอัตโนมัติ']
    if r['fiscal_year']==2567:
        issues.append('บางหมวดในต้นฉบับคำนวณศูนย์จากรายละเอียดว่าง เก็บช่องเหล่านั้นเป็นยังไม่ทราบจนกว่าหน่วยตรวจสอบ')
    else:
        issues.append('รายงานอ้าง 1 ก.ค. 2568 สำหรับรายรับและเงินคงเหลือ ต้องตรวจวันตัดยอดจ่ายจริงแยกต่างหาก')
        issues.append('วงเงินแผนมาจากไฟล์สรุปปี 68 มีฉบับแก้ไขในไฟล์นำเสนอ ต้องยืนยันฉบับที่ใช้จริง')
    for issue in assessment['checks']:
        if issue.get('unit') == r['facility_code'] and r['fiscal_year']==2568:
            if issue['type']=='presentation_total_conflict':
                issues.append('ยอดหน้าสรุป '+issue['summary_amount']+' บาท แต่รายละเอียด '+issue['detail_amount']+' บาท ต้องตรวจหลักฐานและระบุผลตรวจ')
            if issue['type']=='category_total_mismatch':
                issues.append('แผนฉบับนำเสนอ: ยอดหมวดรวม '+issue['sum']+' บาท ไม่ตรงยอดรวม '+issue['reported']+' บาท')
    records.append({'facility_code':r['facility_code'], 'fiscal_year':r['fiscal_year'], 'status':'draft', 'coverage_start':None,'coverage_end':None,
                    'payload':{'groups':groups}, 'source':{'file':Path(src['path']).name,'sheet':r['sheet'],'sha256':src['sha256'],
                    'period_note':'วันตัดยอดรายจ่ายรอหน่วยยืนยัน','issues':issues,'original':r,'plan_versions':plans}, 'review_note':'','reason':'นำเข้าข้อมูลตั้งต้น รอตรวจสอบ','revision':1})

p = Path(sys.argv[1]) if len(sys.argv)>1 else Path(r'D:/68work_James/งานกองสาสุข/แผนเงินบำรุง/ทำแผนเงินบำรุง70/แผนเงินบำรุง70/วิเคราะห์การเงิน2569_8หน่วย_ล่าสุด.xlsx')
wb = openpyxl.load_workbook(p,data_only=True)
formulas = openpyxl.load_workbook(p,data_only=False)
for code, sheet in [('06116','โทกหัวช้าง'),('06117','บ้านฟ่อน'),('06118','ศรีหมวดเกล้า'),('06119','กล้วยแพะ'),('06120','กล้วยม่วง'),('06121','บ้านกาด'),('06122','แม่กืย'),('45030','ศบส.8')]:
    blocks = {}; current = None; originals=[]; declared_total=None
    for row in wb[sheet].iter_rows():
        a,b,c,d,f = [cell.value for cell in row[:5]]
        idx = row[0].row
        if isinstance(a,(int,float)) and a==int(a) and 1<=a<=11:
            current = int(a)
            blocks[current]={'main':(idx,b,c,d,f),'details':[]}
        elif b and 'รวมทั้งสิ้น' in str(b):
            declared_total=num(d);current=None
        elif current is not None and b:
            # Keep headings too; empty amounts remain unknown, never fabricated zero.
            blocks[current]['details'].append((idx,b,c,d,f))
        if current is not None:
            originals.append({'row':idx,'cells':[{'cell':cell.coordinate,'value':cell.value,'formula':formulas[sheet][cell.coordinate].value} for cell in row[:5]]})
    assert len(blocks)==11,(code,'Expected 11 categories')
    groups=[]
    for cat,block in blocks.items():
        details=[{'id':'r'+str(r),'label':str(label)[:500],'amount':num(actual),'commitment':num(committed),'planned':num(plan)} for r,label,plan,actual,committed in block['details']]
        _,label,plan,actual,committed=block['main']
        def chosen(value,field):
            numbers=[d[field] for d in details if d[field] is not None]
            return num(value) if num(value) is not None else round(sum(numbers),2) if numbers else None
        groups.append({'id':str(cat),'label':labels[cat-1],'amount':chosen(actual,'amount'),'commitment':chosen(committed,'commitment'),'planned':chosen(plan,'planned'),'note':'','details':details})
    issues=['จ่ายจริงถึง 31 ส.ค. 2569 เท่านั้น งบผูกพันเดือนกันยายนแยกจากจ่ายจริง','ช่องว่างและเครื่องหมายขีดในต้นฉบับยังไม่ถือเป็นศูนย์ รายละเอียดที่ว่างต้องตรวจสอบก่อนยืนยัน','ยอดรายหมวดต้องเทียบยอดรายเดือนในระบบเฉพาะช่วงเดียวกัน ไม่เติมหรือกระจายส่วนต่างเอง']
    known=round(sum(g['amount'] or 0 for g in groups),2)
    if declared_total is not None and abs(known-declared_total)>.01:
        issues.append('ผลรวมหมวดที่อ่านได้ '+str(known)+' บาท ต่างจากยอดรวมต้นฉบับ '+str(declared_total)+' บาท กรุณาตรวจการจัดหัวข้อ')
    records.append({'facility_code':code,'fiscal_year':2569,'status':'draft','coverage_start':'2025-10-01','coverage_end':'2026-08-31',
                    'payload':{'groups':groups},'source':{'file':p.name,'sheet':sheet,'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),
                    'period_note':'จ่ายจริง 1 ต.ค. 2568–31 ส.ค. 2569 และงบผูกพัน ก.ย. 2569 แยกกัน','issues':issues,'original_rows':originals,'reported_total':declared_total},
                    'review_note':'','reason':'นำเข้าข้อมูลตั้งต้น รอตรวจสอบ','revision':1})

assert len(records)==24 and len({(r['facility_code'],r['fiscal_year']) for r in records})==24
(OUT/'import-records.json').write_text(json.dumps(records,ensure_ascii=False,indent=2,default=str),encoding='utf8')
print(json.dumps({'records':len(records),'years':sorted({r['fiscal_year'] for r in records}),'status':'draft','output':str(OUT/'import-records.json')},ensure_ascii=False))
