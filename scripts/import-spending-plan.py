"""Read-only extraction of the eight FY2570 workbooks; source files are never changed."""
import json, re, sys, hashlib, ast, operator
from pathlib import Path
import openpyxl
ROOT=Path(sys.argv[1]); OUT=Path(sys.argv[2]); OUT.mkdir(parents=True,exist_ok=True)
manifest=[
 ('06116','บ้านโทกหัวช้าง','5.แผนการใช้จ่ายเงินบำรุง ศบส.บ้านโทกหัวช้าง ประจำปีงบประมาณ 2570.xlsx','แผนเงินบำรุง 2570'),
 ('06117','บ้านฟ่อน','1.แผนเงินบำรุง ศบส.บ้านฟ่อน ปีงบประมาณ 2570.xlsx','สรุปแผนปี 70'),
 ('06118','บ้านศรีหมวดเกล้า','แผนเงินบำรุงศรีหมวดเกล้า_ปี2570_แก้ไขใช้อัน.xlsx','สรุปแผน 70 ใช้อันนี้'),
 ('06119','บ้านกล้วยแพะ','แผนเงินบำรุง กล้วยแพะ 2570.xlsx','กล้วยแพะ'),
 ('06120','บ้านกล้วยม่วง','แผนเงินบำรุง ศูนย์ฯ กล้วยม่วง ปีงบประมาณ 2570.xlsx','หน้า 1'),
 ('06121','บ้านกาด','บ้านกาด-แผนการใช้จ่ายเงินบำรุง 2570.xlsx','แผนเงินบำรุง 2570'),
 ('06122','บ้านแม่กืย','แผนการใช้จ่ายเงินบำรุง 2570-ศบส แม่กืย.xlsx','สรุปแผน2570'),
 ('45030','เทศบาลเมืองเขลางค์นคร','แผนเงินบำรุง_ศบส8_ปี2570 ส่งเจมส์180969.xlsx','หน้าแรก')]
plans=[]; report=[]
def number(v):
    if isinstance(v,(int,float)) and not isinstance(v,bool): return round(float(v),2)
    if isinstance(v,str):
        try:return round(float(v.replace(',','').strip()),2)
        except ValueError:pass
    return None
for code,unit,filename,summary_sheet in manifest:
    file=ROOT/filename; wb=openpyxl.load_workbook(file,data_only=True); formulas=openpyxl.load_workbook(file,data_only=False)
    cats={}
    for row in wb[summary_sheet].iter_rows(values_only=True):
        if len(row)>1 and isinstance(row[0],(int,float)) and row[0]==int(row[0]) and 1<=row[0]<=11 and str(int(row[0])) not in cats:
            values=[number(row[i]) for i in (16,17) if i<len(row)]
            value=next((v for v in values if v is not None),None)
            if value is None:
                funding=[number(row[i]) for i in (2,3)]
                value=sum(v for v in funding if v is not None) if any(v is not None for v in funding) else 0
            cats[str(int(row[0]))]={'total':value}
    assert len(cats)==11,(unit,'incomplete summary categories')
    meta={'cats':cats,'overall':round(sum(v['total'] for v in cats.values()),2)}
    # Detail block follows the summary in the two consolidated workbooks.
    ws=wb.worksheets[0] if code in ('06118','06122') else wb.worksheets[1]
    starts=[r for r in range(1,ws.max_row+1) if ws.cell(r,1).value==1 and str(ws.cell(r,2).value or '').startswith('ค่ายา')]
    start=starts[-1]; groups=[]; group=None
    for r in range(start,ws.max_row+1):
        a=ws.cell(r,1).value; label=str(ws.cell(r,2).value or '').strip()
        if isinstance(a,(int,float)) and a==int(a) and 1<=a<=11:
            gid=str(int(a)); group={'id':gid,'label':label,'annual':meta['cats'][gid]['total'],'rows':[]};groups.append(group)
        elif label=='รวม' and group and group['id']=='11':break
        elif not group or not label or label in ('รายการ','อปท','เงินบำรุง') or a=='ลำดับที่':continue
        if not group:continue
        is_main=isinstance(a,(int,float)) and a==int(a) and 1<=a<=11
        raw={openpyxl.utils.get_column_letter(c):ws.cell(r,c).value for c in range(1,ws.max_column+1) if ws.cell(r,c).value is not None}
        raw_formula={openpyxl.utils.get_column_letter(c):formulas[ws.title].cell(r,c).value for c in range(1,ws.max_column+1) if formulas[ws.title].cell(r,c).data_type=='f'}
        # R is the corrected annual total column in Mae Kuey's detailed plan.
        annual=number(ws.cell(r,18 if code=='06122' else 17).value)
        months=[number(ws.cell(r,c).value) for c in range(5,17)]
        # Recalculate simple arithmetic and same-sheet SUM only when caches are absent.
        def calc(addr,seen=None):
            seen=set() if seen is None else set(seen)
            if addr in seen:return None
            seen.add(addr); cell=formulas[ws.title][addr];v=number(wb[ws.title][addr].value)
            if v is not None:return v
            f=cell.value
            if not isinstance(f,str) or not f.startswith('=') or '#REF!' in f:return None
            expr=f[1:].replace('$','')
            def summ(m):
                vals=[calc(c.coordinate,seen) for row in formulas[ws.title][m[1]] for c in row]
                if any(x is None and formulas[ws.title][c.coordinate].data_type=='f' for row in formulas[ws.title][m[1]] for c,x in zip(row,[calc(c.coordinate,seen) for c in row])):raise ValueError()
                return str(sum(x or 0 for x in vals))
            try:
                expr=re.sub(r'SUM\(([A-Z]+\d+:[A-Z]+\d+)\)',summ,expr,flags=re.I)
                def ref(m):
                    v=calc(m[0],seen)
                    if v is None and formulas[ws.title][m[0]].data_type=='f':raise ValueError()
                    return str(v or 0)
                expr=re.sub(r'\b[A-Z]+\d+\b',ref,expr)
                if not re.fullmatch(r'[\d\s.+*/()\-]+',expr):return None
                tree=ast.parse(expr,mode='eval')
                ops={ast.Add:operator.add,ast.Sub:operator.sub,ast.Mult:operator.mul,ast.Div:operator.truediv}
                def ev(n):
                    if isinstance(n,ast.Constant) and isinstance(n.value,(int,float)):return n.value
                    if isinstance(n,ast.BinOp) and type(n.op) in ops:return ops[type(n.op)](ev(n.left),ev(n.right))
                    if isinstance(n,ast.UnaryOp) and isinstance(n.op,(ast.USub,ast.UAdd)):return (-1 if isinstance(n.op,ast.USub) else 1)*ev(n.operand)
                    raise ValueError()
                return round(ev(tree.body),2)
            except (ValueError,TypeError,ZeroDivisionError,SyntaxError):return None
        if annual is None:annual=calc(('R' if code=='06122' else 'Q')+str(r))
        for i in range(12):
            if months[i] is None:months[i]=calc(openpyxl.utils.get_column_letter(i+5)+str(r))
        match=re.match(r'^(\d+(?:\.\d+)+)(?=\D|$)',label)
        group['rows'].append({'id':'r'+str(r),'code':match[1] if match else group['id'] if is_main else '', 'label':label,'annual':annual,'months':months,'source_row':r,'source_cells':raw,'source_formulas':raw_formula,'main':is_main})
    for g in groups:
        parent=g['id']
        for row in g['rows']:
            if row['code']:parent=row['code']
            elif row['annual'] is not None and row['annual']>0 and not row['label'].startswith(('ครุภัณฑ์','ที่ดินและสิ่งก่อสร้าง','ค่าครุภัณฑ์')):
                # Some workbooks list staff/service subitems as (1), (2), or plain text.
                row['code']=parent+'.u'+row['id'];row['unnumbered']=True
        for i,row in enumerate(g['rows']):
            following=g['rows'][i+1:]
            children=[x for x in following if x['code'].startswith(row['code']+'.')] if row['code'] else []
            row['input']=bool(row['code']) and not children
            leaf_children=[x for x in children if not any(y['code'].startswith(x['code']+'.') for y in children)]
            if children and not row['main'] and row['annual'] and all(x['annual'] is not None for x in leaf_children):
                same_total=abs(sum(x['annual'] for x in leaf_children)-row['annual'])<0.01
                formula_refs=[int(n) for f in row['source_formulas'].values() for n in re.findall(r'[A-Z]+(\d+)',str(f))]
                sums_children=any(x['source_row'] in formula_refs for x in children)
                if not same_total and not sums_children:
                    row['input']=True;row['separate_budget']=True
            if row['main']:row['input']=len(g['rows'])==1
        # Preserve unnumbered descriptive headings without making them duplicate expense inputs.
    assert len(groups)==11,(code,len(groups))
    issues=[]
    for g in groups:
        leaves=[r for r in g['rows'] if r['input']]
        if any(r['annual'] is None for r in leaves):issues.append({'group':g['id'],'kind':'unallocated','message':'ต้นฉบับไม่ได้แจกแจงวงเงินให้ทุกรายการย่อย ใช้วงเงินรวมของหมวดในการติดตาม'})
        elif abs(sum(r['annual'] for r in leaves)-g['annual'])>=0.01:issues.append({'group':g['id'],'kind':'difference','message':'ผลรวมวงเงินรายการย่อยต่างจากแผนสรุปต้นฉบับ','difference':round(sum(r['annual'] for r in leaves)-g['annual'],2)})
    payload={'source_file':file.name,'source_sheet':ws.title,'source_sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'groups':groups,'annual':meta['overall'],'source_issues':issues}
    plans.append({'facility_code':code,'fiscal_year':2570,'period':0,'status':'plan','payload':payload})
    report.append({'code':code,'unit':unit,'groups':len(groups),'rows':sum(len(g['rows']) for g in groups),'inputs':sum(r['input'] for g in groups for r in g['rows']),'annual':meta['overall'],'leaf_annual':sum(r['annual'] or 0 for g in groups for r in g['rows'] if r['input'])})
(OUT/'spending-plans-2570.json').write_text(json.dumps(plans,ensure_ascii=False,default=str),encoding='utf-8')
(OUT/'spending-plan-import-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
